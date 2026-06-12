#!/usr/bin/env python3
import os
import re
import json
import time
import random
import argparse
import urllib.parse
from datetime import datetime, date, timedelta
from pathlib import Path
import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

COOKIE_ROTATE_INTERVAL = 8

MOBILE_USER_AGENTS = [
    "Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) AppleWebKit/605.1.15 Mobile/15E148 Safari/604.1",
    "Mozilla/5.0 (iPhone; CPU iPhone OS 16_5 like Mac OS X) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/16.5 Mobile/15E148 Safari/604.1",
    "Mozilla/5.0 (iPhone; CPU iPhone OS 15_7 like Mac OS X) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/15.7 Mobile/15E148 Safari/604.1",
    "Mozilla/5.0 (Linux; Android 13; SM-S901B) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/112.0.0.0 Mobile Safari/537.36",
    "Mozilla/5.0 (Linux; Android 12; Pixel 6) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/112.0.0.0 Mobile Safari/537.36",
    "Mozilla/5.0 (Linux; Android 11; SM-A515F) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/111.0.0.0 Mobile Safari/537.36",
    "Mozilla/5.0 (Linux; Android 10; Redmi Note 9 Pro) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/110.0.0.0 Mobile Safari/537.36"
]

def get_randomized_headers(referer: str = "https://m.weibo.cn/") -> dict:
    ua = random.choice(MOBILE_USER_AGENTS)
    rand_ip = f"{random.randint(1, 254)}.{random.randint(1, 254)}.{random.randint(1, 254)}.{random.randint(1, 254)}"
    return {
        "User-Agent": ua,
        "Accept": "application/json,text/plain,*/*",
        "Accept-Language": random.choice(["zh-CN,zh;q=0.9,en;q=0.8", "zh-CN,zh;q=0.9", "zh-Hans;q=0.9,en-US;q=0.8"]),
        "Referer": referer,
        "X-Forwarded-For": rand_ip,
        "Client-IP": rand_ip,
        "X-Real-IP": rand_ip,
        "X-Client-IP": rand_ip
    }

class CrawlerConfig:
    def __init__(self, config_path: Path):
        self.project_root = config_path.parent
        self.keyword = "尼帕病毒"
        self.start_date_str = "2026-01-15"
        self.end_date_str = "2026-03-15"
        self.max_pages = 90
        self.output_dir_name = "outputs"
        
        if config_path.exists():
            try:
                data = json.loads(config_path.read_text(encoding="utf-8"))
                self.keyword = data.get("keyword", self.keyword)
                self.start_date_str = data.get("start_date", self.start_date_str)
                self.end_date_str = data.get("end_date", self.end_date_str)
                self.max_pages = int(data.get("max_pages", self.max_pages))
                self.output_dir_name = data.get("output_dir", self.output_dir_name)
            except Exception as e:
                print(f"Warning: Failed to load config: {e}. Using default settings.", flush=True)
                
        self.start_date = datetime.strptime(self.start_date_str, "%Y-%m-%d").date()
        self.end_date = datetime.strptime(self.end_date_str, "%Y-%m-%d").date()
        
        self.output_dir = self.project_root / self.output_dir_name
        self.output_dir.mkdir(parents=True, exist_ok=True)
        
        self.raw_path = self.output_dir / "search_mblogs.jsonl"
        self.detail_cache_path = self.output_dir / "detail_cache.json"
        self.progress_path = self.output_dir / "progress.json"
        self.cookies_path = self.output_dir / "weibo_mobile_cookies.json"
        self.output_xlsx = self.output_dir / f"微博_{self.keyword}_数据集_{self.start_date_str}_{self.end_date_str}_单机版.xlsx"

class LocalSession:
    def __init__(self, cookies_path: Path):
        self.session = requests.Session()
        self.cookies_path = cookies_path
        self.cookie = {}
        self.request_count = 0
        
        if self.cookies_path.exists():
            try:
                loaded = json.loads(self.cookies_path.read_text(encoding="utf-8"))
                # Unpack nested worker pool cookies if present
                if isinstance(loaded, dict):
                    if any(isinstance(val, dict) for val in loaded.values()):
                        for val in loaded.values():
                            if isinstance(val, dict) and "SUB" in val:
                                self.cookie = val
                                break
                    else:
                        self.cookie = loaded
            except Exception:
                pass
        self._apply_cookies()

    def _apply_cookies(self):
        self.session.cookies.clear()
        for k, v in self.cookie.items():
            self.session.cookies.set(k, v, domain=".weibo.cn")

    def bootstrap(self, max_retries=3):
        search_url = "https://m.weibo.cn/search?containerid=100103type%3D1%26q%3D%E5%B0%BC%E5%B8%95%E7%97%85%E6%AF%92"
        headers = {
            "User-Agent": "Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) AppleWebKit/605.1.15 Mobile/15E148 Safari/604.1",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "zh-CN,zh;q=0.9",
        }
        for attempt in range(1, max_retries + 1):
            try:
                print(f"[Session] 自动引导新浪访客 Cookie (第 {attempt}/{max_retries} 次尝试)...", flush=True)
                resp = self.session.get(search_url, headers=headers, timeout=20)
                if resp.status_code != 200:
                    continue
                req_id_match = re.search(r'var request_id = "([^"]+)"', resp.text)
                return_url_match = re.search(r'var return_url = "([^"]+)"', resp.text)
                if not req_id_match or not return_url_match:
                    continue
                
                req_id = req_id_match.group(1)
                return_url = return_url_match.group(1)
                post_body = f"cb=visitor_gray_callback&ver=20250916&request_id={req_id}&tid=&from=weibo&webdriver=undefined&rid={int(time.time()*1000)}&return_url={urllib.parse.quote(return_url)}"
                
                headers2 = {
                    "Origin": "https://visitor.passport.weibo.cn",
                    "Referer": "https://visitor.passport.weibo.cn/visitor/visitor",
                    "Content-Type": "application/x-www-form-urlencoded",
                    "User-Agent": headers["User-Agent"]
                }
                
                resp2 = self.session.post("https://visitor.passport.weibo.cn/visitor/genvisitor2", data=post_body, headers=headers2, timeout=20)
                if resp2.status_code != 200:
                    continue
                match = re.search(r'visitor_gray_callback\((.*?)\);', resp2.text)
                if not match:
                    continue
                
                res_data = json.loads(match.group(1))
                if res_data.get("retcode") != 20000000:
                    continue
                
                self.cookie = {
                    "SUB": res_data["data"]["sub"],
                    "SUBP": res_data["data"]["subp"],
                    "WEIBOCN_FROM": "1110003030"
                }
                self._apply_cookies()
                self.request_count = 0
                
                self.cookies_path.write_text(json.dumps(self.cookie, ensure_ascii=False, indent=2), encoding="utf-8")
                print("[Session] 访客 Cookie 引导成功！", flush=True)
                return True
            except Exception as e:
                print(f"[Session] 引导失败: {e}", flush=True)
                time.sleep(1.5)
        return False

    def request(self, url, method="GET", headers=None, params=None, timeout=20):
        self.request_count += 1
        if self.request_count > COOKIE_ROTATE_INTERVAL:
            print(f"[Session] 请求量 {self.request_count} > {COOKIE_ROTATE_INTERVAL}，自动重置并刷新 Cookie...", flush=True)
            self.cookie = {}
            self.request_count = 0
            
        if not self.cookie:
            self.bootstrap()
            
        req_headers = get_randomized_headers(url)
        if headers:
            req_headers.update(headers)
            
        for attempt in range(2):
            try:
                if method == "GET":
                    resp = self.session.get(url, headers=req_headers, params=params, timeout=timeout)
                else:
                    resp = self.session.post(url, headers=req_headers, params=params, timeout=timeout)
                    
                is_captcha = "captcha/show" in resp.text or "微博验证" in resp.text or "Sina Visitor System" in resp.text or resp.status_code in (401, 403, 418)
                if is_captcha:
                    print(f"[Session] 遇到验证码或频次限制。重新引导 Cookie 身份...", flush=True)
                    self.cookie = {}
                    self.request_count = 0
                    if self.bootstrap():
                        self._apply_cookies()
                        continue
                    else:
                        break
                return resp
            except Exception as e:
                print(f"[Session] 请求异常: {e}", flush=True)
                time.sleep(1.5)
        raise RuntimeError("网络连接失败，请检查网络后重试。")

def clean_text(value: str) -> str:
    import html
    value = re.sub(r"<br\s*/?>", "\n", value or "", flags=re.I)
    value = re.sub(r"<[^>]+>", "", value)
    return html.unescape(value).replace("\u200b", "").strip()

def parse_created_at(value: str, today: date = None) -> datetime | None:
    today = today or date.today()
    value = (value or "").strip()
    if not value: return None
    now = datetime.now()
    if "刚刚" in value: return now
    minute = re.search(r"(\d+)\s*分钟前", value)
    if minute: return now - timedelta(minutes=int(minute.group(1)))
    hour = re.search(r"(\d+)\s*小时前", value)
    if hour: return now - timedelta(hours=int(hour.group(1)))
    if value.startswith("今天"): return datetime.combine(today, datetime.min.time())
    if value.startswith("昨天"): return datetime.combine(today - timedelta(days=1), datetime.min.time())
    for fmt in ("%a %b %d %H:%M:%S %z %Y", "%a %b %d %H:%M:%S %Y"):
        try:
            return datetime.strptime(value, fmt).replace(tzinfo=None)
        except ValueError:
            pass
    match = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:\s+(\d{1,2}):(\d{1,2}))?", value)
    if match:
        year, month, day, hour_s, minute_s = match.groups()
        return datetime(int(year), int(month), int(day), int(hour_s or 0), int(minute_s or 0))
    if re.fullmatch(r"\d{1,2}-\d{1,2}(?:\s+\d{1,2}:\d{1,2})?", value):
        left, *time_part = value.split()
        month, day = map(int, left.split("-"))
        hour_i = minute_i = 0
        if time_part:
            hour_i, minute_i = map(int, time_part[0].split(":"))
        return datetime(today.year, month, day, hour_i, minute_i)
    return None

def collect_mblogs(cards: list, out: list = None) -> list:
    out = out if out is not None else []
    for card in cards or []:
        mblog = card.get("mblog")
        if isinstance(mblog, dict):
            out.append(mblog)
        group = card.get("card_group")
        if isinstance(group, list):
            collect_mblogs(group, out)
    return out

def fetch_search_page(session, keyword: str, page: int, search_type: int = 1) -> dict:
    url = "https://m.weibo.cn/api/container/getIndex"
    params = {
        "containerid": f"100103type={search_type}&q={keyword}",
        "page": page
    }
    headers = {
        "Accept": "application/json, text/plain, */*",
        "Referer": f"https://m.weibo.cn/search?containerid=100103type%3D{search_type}%26q%3D{urllib.parse.quote(keyword)}"
    }
    resp = session.request(url, method="GET", headers=headers, params=params)
    if resp.status_code != 200:
        return {"ok": 0, "status": resp.status_code}
    try:
        return resp.json()
    except Exception as e:
        return {"ok": 0, "error": str(e)}

def fetch_detail(session, bid: str) -> dict:
    url = f"https://m.weibo.cn/detail/{bid}"
    headers = {
        "Accept": "text/html,application/xhtml+xml,application/json,text/plain,*/*",
        "Referer": "https://m.weibo.cn/"
    }
    try:
        resp = session.request(url, method="GET", headers=headers)
        if resp.status_code != 200:
            return {"status": f"http_{resp.status_code}", "mblog": None}
        match = re.search(r"var \$render_data = (\[.*?\])\[0\]", resp.text, re.S)
        if not match:
            return {"status": "render_missing", "mblog": None}
        render = json.loads(match.group(1))
        status = (render[0] or {}).get("status") if render else None
        if isinstance(status, dict):
            return {"status": "ok", "mblog": status}
        return {"status": "empty_status", "mblog": None}
    except Exception as e:
        return {"status": f"exception_{e}", "mblog": None}

def build_excel(config, posts: list):
    wb = Workbook()
    thin = Side(style="thin", color="D9D9D9")
    fill = PatternFill("solid", fgColor="4472C4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    
    # Sheet 1
    ws_posts = wb.active
    ws_posts.title = "微博帖子"
    post_headers = [
        "平台", "关键词", "微博ID", "短ID", "博主昵称", "用户ID", "发布时间", "发布日期", "微博内容",
        "话题标签", "IP/发布地区", "转发数", "评论数", "点赞数", "发布来源", "是否长文",
        "图片数量", "图片链接", "视频链接", "微博链接", "详情抓取状态"
    ]
    ws_posts.append(post_headers)
    
    sorted_posts = sorted(posts, key=lambda p: str(p.get("created_at") or ""), reverse=True)
    
    for p in sorted_posts:
        user = p.get("user") or {}
        bid = str(p.get("bid") or p.get("mblogid") or p.get("id") or "")
        created_time = str(p.get("created_at") or "")
        dt = parse_created_at(created_time)
        created_time_str = dt.strftime("%Y-%m-%d %H:%M:%S") if dt else created_time
        created_date_str = dt.strftime("%Y-%m-%d") if dt else ""
        text = clean_text(str(p.get("text_raw") or p.get("text") or ""))
        
        pics = []
        for pic in p.get("pics") or []:
            url = (pic.get("large") or {}).get("url") or pic.get("url")
            if url: pics.append(url)
            
        page_info = p.get("page_info") or {}
        media = page_info.get("media_info") or {}
        v_url = media.get("stream_url") or media.get("mp4_hd_url") or media.get("mp4_sd_url") or ""
        
        tags = re.findall(r"#([^#]{1,80})#", text)
        tags_str = "、".join(dict.fromkeys(tag.strip() for tag in tags if tag.strip()))
        
        region = p.get("region_name") or ""
        province = p.get("status_province") or ""
        city = p.get("status_city") or ""
        if not region and (province or city):
            region = " ".join(str(x) for x in [province, city] if x)
            
        ws_posts.append([
            "微博", p.get("matched_keyword", config.keyword), str(p.get("id") or ""), bid,
            user.get("screen_name") or "", str(user.get("id") or ""), created_time_str, created_date_str,
            text, tags_str, region, int(p.get("reposts_count") or 0), int(p.get("comments_count") or 0),
            int(p.get("attitudes_count") or 0), clean_text(str(p.get("source") or "")),
            "是" if p.get("isLongText") else "否", len(pics), " ".join(pics), v_url,
            f"https://weibo.com/{user.get('id') or ''}/{bid}" if bid else "", "ok"
        ])

    for cell in ws_posts[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    widths_posts = {
        "A": 8, "B": 12, "C": 20, "D": 14, "E": 18, "F": 16, "G": 20, "H": 14, "I": 90,
        "J": 28, "K": 18, "L": 10, "M": 10, "N": 10, "O": 24, "P": 10, "Q": 10,
        "R": 50, "S": 40, "T": 42, "U": 16
    }
    for col, width in widths_posts.items():
        ws_posts.column_dimensions[col].width = width
    for row in ws_posts.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws_posts.freeze_panes = "A2"
    ws_posts.auto_filter.ref = ws_posts.dimensions

    # Sheet 2
    ws_summary = wb.create_sheet("统计说明")
    summary_rows = [
        ["项目", "数量/内容"],
        ["关键词", config.keyword],
        ["时间范围", f"{config.start_date_str} 至 {config.end_date_str}"],
        ["去重后命中帖子数", len(posts)],
        ["说明", "通过微博移动端检索接口翻页，获取相关的微博帖子及详情（包括全文和发布地区/IP）。"]
    ]
    for row in summary_rows:
        ws_summary.append(row)
    for cell in ws_summary[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws_summary.column_dimensions["A"].width = 24
    ws_summary.column_dimensions["B"].width = 80
    for row in ws_summary.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(config.output_xlsx)
    print(f"📊 最终 Excel 报表已生成: {config.output_xlsx}", flush=True)

def main():
    parser = argparse.ArgumentParser(description="微博单机版关键词爬虫工具 (不带代理Worker，适合本地独立运行)")
    parser.add_argument("--config", type=str, default="config.json", help="配置文件路径")
    parser.add_argument("--skip-search", action="store_true", help="跳过关键词翻页抓取步骤")
    args = parser.parse_args()
    
    project_dir = Path(__file__).resolve().parent
    config = CrawlerConfig(project_dir / args.config)
    
    print("=" * 60)
    print(" 🚀 微博关键词爬虫 (单机版) 启动 ")
    print(f" 📂 工作根目录: {config.project_root}")
    print(f" 🔑 检索关键词: {config.keyword}")
    print(f" 📅 时间过滤范围: {config.start_date_str} 至 {config.end_date_str}")
    print(f" 📊 输出 Excel 路径: {config.output_xlsx}")
    print("=" * 60, flush=True)

    session = LocalSession(config.cookies_path)

    # Step 1: Search posts
    if not args.skip_search:
        print("开始抓取主帖...", flush=True)
        progress = {}
        if config.progress_path.exists():
            try: progress = json.loads(config.progress_path.read_text(encoding="utf-8"))
            except Exception: pass

        seen_bids = set()
        if config.raw_path.exists():
            with config.raw_path.open("r", encoding="utf-8") as f:
                for line in f:
                    try:
                        item = json.loads(line)
                        if item.get("bid"): seen_bids.add(item["bid"])
                    except Exception: pass

        for st in [1, 61]:
            task_key = f"{config.keyword}_{st}_simple"
            if progress.get(f"{task_key}_completed"):
                print(f"Task '{task_key}' already completed, skipping.", flush=True)
                continue
                
            start_page = progress.get(f"{task_key}_last_page", 0) + 1
            print(f"正在进行搜素 type={st} 从第 {start_page} 页开始...", flush=True)
            
            completed = True
            for page in range(start_page, config.max_pages + 1):
                try:
                    payload = fetch_search_page(session, config.keyword, page, search_type=st)
                    if payload.get("ok") != 1:
                        print(f"已到达翻页尽头，或被安全拦截。结束当前搜索。", flush=True)
                        progress[f"{task_key}_last_page"] = page - 1
                        break
                    
                    cards = (payload.get("data") or {}).get("cards") or []
                    mblogs = collect_mblogs(cards)
                    print(f"第 {page} 页: 找到 {len(mblogs)} 条帖子", flush=True)
                    
                    if not mblogs:
                        progress[f"{task_key}_last_page"] = page
                        break
                        
                    new_found = 0
                    with config.raw_path.open("a", encoding="utf-8") as f:
                        for mblog in mblogs:
                            bid = str(mblog.get("bid") or mblog.get("mblogid") or mblog.get("id") or "")
                            if not bid or bid in seen_bids: continue
                            seen_bids.add(bid)
                            new_found += 1
                            mblog["matched_keyword"] = config.keyword
                            f.write(json.dumps({"bid": bid, "mblog": mblog}, ensure_ascii=False) + "\n")
                    
                    progress[f"{task_key}_last_page"] = page
                    config.progress_path.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
                    time.sleep(random.uniform(2.5, 4.5)) # 单机版请求增加延迟防止被限制
                except Exception as e:
                    print(f"抓取主帖异常: {e}", flush=True)
                    completed = False
                    progress[f"{task_key}_last_page"] = page - 1
                    break
            if completed:
                progress[f"{task_key}_completed"] = True
                config.progress_path.write_text(json.dumps(progress, ensure_ascii=False, indent=2), encoding="utf-8")
    else:
        print("跳过主帖抓取，直接加载本地缓存...", flush=True)

    # Step 2: Detail enrichment
    detail_cache = {}
    if config.detail_cache_path.exists():
        try: detail_cache = json.loads(config.detail_cache_path.read_text(encoding="utf-8"))
        except Exception: pass

    all_raw_posts = []
    seen_bids = set()
    if config.raw_path.exists():
        with config.raw_path.open("r", encoding="utf-8") as f:
            for line in f:
                try:
                    item = json.loads(line)
                    mblog = item.get("mblog")
                    bid = item.get("bid")
                    if mblog and bid and bid not in seen_bids:
                        created = mblog.get("created_at", "")
                        dt = parse_created_at(created)
                        if dt and config.start_date <= dt.date() <= config.end_date:
                            seen_bids.add(bid)
                            all_raw_posts.append(mblog)
                except Exception: pass

    print(f"开始补全帖子详情... 待处理目标帖子总量: {len(all_raw_posts)}", flush=True)
    detail_counter = 0
    for idx, post in enumerate(all_raw_posts, start=1):
        bid = str(post.get("bid") or post.get("mblogid") or post.get("id") or "")
        cached = detail_cache.get(bid)
        if cached is None or cached.get("status") not in ("ok", "cached"):
            res = fetch_detail(session, bid)
            detail_cache[bid] = {"status": res.get("status"), "mblog": res.get("mblog") or post}
            detail_counter += 1
            if detail_counter % 5 == 0:
                config.detail_cache_path.write_text(json.dumps(detail_cache, ensure_ascii=False, indent=2), encoding="utf-8")
            print(f"[{idx}/{len(all_raw_posts)}] 补全详情 bid={bid} | 状态: {res.get('status')}", flush=True)
            time.sleep(random.uniform(1.2, 2.5))
    config.detail_cache_path.write_text(json.dumps(detail_cache, ensure_ascii=False, indent=2), encoding="utf-8")
    print("详情补全完毕！", flush=True)

    valid_posts = []
    for bid, item in detail_cache.items():
        mblog = item.get("mblog")
        if item.get("status") in ("ok", "cached") and isinstance(mblog, dict):
            created = mblog.get("created_at", "")
            dt = parse_created_at(created)
            if dt and config.start_date <= dt.date() <= config.end_date:
                valid_posts.append(mblog)

    # Step 3: Export to Excel (Skip comments fetching completely)
    build_excel(config, valid_posts)
    print("=" * 60)
    print(" 🎉 流程全部运行完毕！(单机版)")
    print(f" 📂 报表文件已输出至: {config.output_xlsx}")
    print("=" * 60)

if __name__ == "__main__":
    main()
