import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from .search_crawler import clean_text

def parse_created_at_local(value: str) -> str:
    from datetime import datetime
    for fmt in ("%a %b %d %H:%M:%S %z %Y", "%a %b %d %H:%M:%S %Y"):
        try:
            parsed = datetime.strptime(value, fmt)
            return parsed.strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            pass
    match = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:\s+(\d{1,2}):(\d{1,2}))?", value)
    if match:
        year, month, day, hour_s, minute_s = match.groups()
        return f"{int(year):04d}-{int(month):02d}-{int(day):02d} {int(hour_s or 0):02d}:{int(minute_s or 0):02d}:00"
    return value

def pic_urls(mblog: dict) -> list:
    urls = []
    for pic in mblog.get("pics") or []:
        url = (pic.get("large") or {}).get("url") or pic.get("url")
        if url:
            urls.append(url)
    return urls

def video_url(mblog: dict) -> str:
    page_info = mblog.get("page_info") or {}
    media = page_info.get("media_info") or {}
    return media.get("stream_url") or media.get("mp4_hd_url") or media.get("mp4_sd_url") or ""

def tags_from_text(value: str) -> str:
    tags = re.findall(r"#([^#]{1,80})#", clean_text(value))
    return "、".join(dict.fromkeys(tag.strip() for tag in tags if tag.strip()))

def build_excel_workbook(config, posts: list, comments_cache: dict):
    print("Building multi-sheet Excel workbook...", flush=True)
    wb = Workbook()
    thin = Side(style="thin", color="D9D9D9")
    fill = PatternFill("solid", fgColor="4472C4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    
    # ------------------ SHEET 1: 微博帖子 ------------------
    ws_posts = wb.active
    ws_posts.title = "微博帖子"
    post_headers = [
        "平台", "关键词", "微博ID", "短ID", "博主昵称", "用户ID", "发布时间", "发布日期", "微博内容",
        "话题标签", "IP/发布地区", "转发数", "评论数", "点赞数", "发布来源", "是否长文",
        "图片数量", "图片链接", "视频链接", "微博链接", "详情抓取状态"
    ]
    ws_posts.append(post_headers)
    
    # Sort posts by created_at desc
    sorted_posts = sorted(posts, key=lambda p: parse_created_at_local(str(p.get("created_at") or "")), reverse=True)
    
    for p in sorted_posts:
        user = p.get("user") or {}
        bid = str(p.get("bid") or p.get("mblogid") or p.get("id") or "")
        created_time = parse_created_at_local(str(p.get("created_at") or ""))
        created_date = created_time.split()[0] if created_time and " " in created_time else ""
        text = clean_text(str(p.get("text_raw") or p.get("text") or ""))
        images = pic_urls(p)
        region = p.get("region_name") or ""
        province = p.get("status_province") or ""
        city = p.get("status_city") or ""
        if not region and (province or city):
            region = " ".join(str(x) for x in [province, city] if x)
            
        row = [
            "微博",
            p.get("matched_keyword", config.keyword),
            str(p.get("id") or ""),
            bid,
            user.get("screen_name") or "",
            str(user.get("id") or ""),
            created_time,
            created_date,
            text,
            tags_from_text(text),
            region,
            int(p.get("reposts_count") or 0),
            int(p.get("comments_count") or 0),
            int(p.get("attitudes_count") or 0),
            clean_text(str(p.get("source") or "")),
            "是" if p.get("isLongText") else "否",
            len(images),
            " ".join(images),
            video_url(p),
            f"https://weibo.com/{user.get('id') or ''}/{bid}" if bid else "",
            "ok"
        ]
        ws_posts.append(row)

    # Style Sheet 1
    for cell in ws_posts[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    widths_posts = {
        "A": 8, "B": 12, "C": 20, "D": 14, "E": 18, "F": 16, "G": 20, "H": 14, "I": 90,
        "J": 28, "K": 18, "L": 10, "M": 10, "N": 10, "O": 24, "P": 10, "Q": 10,
        "R": 50, "S": 40, "T": 42, "U": 16
    }
    for col, width in widths_posts.items():
        ws_posts.column_dimensions[col].width = width
    for row in ws_posts.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws_posts.freeze_panes = "A2"
    ws_posts.auto_filter.ref = ws_posts.dimensions

    # ------------------ SHEET 2: 微博评论 ------------------
    ws_comments = wb.create_sheet("微博评论")
    comment_headers = [
        "平台", "关键词", "原微博ID", "原微博短ID", "原微博内容摘要",
        "评论ID", "评论人昵称", "评论人ID", "评论时间", "评论内容",
        "点赞数", "IP/发布地区", "原微博链接"
    ]
    ws_comments.append(comment_headers)

    for p in sorted_posts:
        bid = str(p.get("bid") or p.get("mblogid") or p.get("id") or "")
        user = p.get("user") or {}
        weibo_id = str(p.get("id") or "")
        keyword = p.get("matched_keyword", config.keyword)
        raw_text = clean_text(str(p.get("text_raw") or p.get("text") or ""))
        text_summary = raw_text[:50] + "..." if len(raw_text) > 50 else raw_text
        weibo_link = f"https://weibo.com/{user.get('id') or ''}/{bid}" if bid else ""
        
        post_comments = comments_cache.get(bid, [])[:config.max_comments]
        for c in post_comments:
            row_cmt = [
                "微博",
                keyword,
                weibo_id,
                bid,
                text_summary,
                c.get("id", ""),
                c.get("user_nickname", ""),
                c.get("user_id", ""),
                c.get("created_at", ""),
                c.get("text", ""),
                c.get("like_count", 0),
                c.get("source", ""),  # IP/发布地区
                weibo_link
            ]
            ws_comments.append(row_cmt)

    # Style Sheet 2
    for cell in ws_comments[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    widths_comments = {
        "A": 8, "B": 12, "C": 20, "D": 14, "E": 36, "F": 20, "G": 18, "H": 16,
        "I": 20, "J": 80, "K": 10, "L": 18, "M": 42
    }
    for col, width in widths_comments.items():
        ws_comments.column_dimensions[col].width = width
    for row in ws_comments.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws_comments.freeze_panes = "A2"
    ws_comments.auto_filter.ref = ws_comments.dimensions

    # ------------------ SHEET 3: 统计说明 ------------------
    ws_summary = wb.create_sheet("统计说明")
    total_comments_count = sum(len(comments_cache.get(str(p.get("bid") or p.get("mblogid") or p.get("id")), [])[:config.max_comments]) for p in posts)
    
    summary_rows = [
        ["项目", "数量/内容"],
        ["关键词", keyword],
        ["时间范围", f"{config.start_date_str} 至 {config.end_date_str}"],
        ["去重后命中帖子数", len(posts)],
        ["成功抓取评论总数", total_comments_count],
        ["单篇帖子抓取评论上限", f"前 {config.max_comments} 条热门一级评论"],
        ["说明", f"通过微博移动端评论热评接口 (/comments/hotflow) 翻页，每个帖子前 {config.max_comments} 条热门评论，含评论发布地区/IP。"]
    ]
    for row in summary_rows:
        ws_summary.append(row)
    for cell in ws_summary[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    ws_summary.column_dimensions["A"].width = 24
    ws_summary.column_dimensions["B"].width = 80
    for row in ws_summary.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Save
    wb.save(config.output_xlsx)
    print(f"Excel workbook built successfully: {config.output_xlsx}", flush=True)
    print(f"Summary -> Posts: {len(posts)} | Comments: {total_comments_count}", flush=True)
